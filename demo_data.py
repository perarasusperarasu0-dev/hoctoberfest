"""
LazySort - Demo Data Generator
Creates realistic sample files (PDF, DOCX, TXT, PNG, JSON) across
Finance, Work, Education, Documents, Images, and Others for safe testing.
"""

from pathlib import Path
from typing import List


def create_sample_pdf(filepath: Path, title: str, paragraphs: List[str]):
    """Creates an authentic PDF document using reportlab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        doc = SimpleDocTemplate(str(filepath), pagesize=letter, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=14
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=11,
            leading=16,
            textColor=colors.HexColor('#334155'),
            spaceAfter=10
        )

        elements = [Paragraph(title, title_style), Spacer(1, 12)]
        for p in paragraphs:
            elements.append(Paragraph(p, body_style))
            elements.append(Spacer(1, 6))

        doc.build(elements)
    except Exception as e:
        # Fallback to plain text if reportlab encounters an unexpected issue
        with open(filepath.with_suffix(".txt"), "w", encoding="utf-8") as f:
            f.write(f"{title}\n\n" + "\n\n".join(paragraphs))


def create_sample_docx(filepath: Path, title: str, sections: List[tuple]):
    """Creates an authentic DOCX Word document using python-docx."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    
    # Title
    head = doc.add_heading(title, level=0)
    head.style.font.name = 'Calibri'
    head.style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    for sec_title, sec_content in sections:
        h = doc.add_heading(sec_title, level=1)
        h.style.font.name = 'Calibri'
        h.style.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        p = doc.add_paragraph(sec_content)
        p.style.font.name = 'Calibri'
        p.style.font.size = Pt(11)

    doc.save(str(filepath))


def create_sample_png(filepath: Path, title: str):
    """Generates a clean illustrative diagram PNG using Pillow."""
    from PIL import Image, ImageDraw, ImageFont

    width, height = 700, 450
    image = Image.new("RGB", (width, height), color=(248, 250, 252))
    draw = ImageDraw.Draw(image)

    # Draw header banner
    draw.rectangle([0, 0, width, 70], fill=(30, 41, 59))
    draw.text((30, 25), title, fill=(255, 255, 255))

    # Draw decorative architecture diagram boxes
    box_coords = [
        ((60, 140, 220, 240), (224, 231, 255), (37, 99, 235), "Web Client\n(Streamlit UI)"),
        ((270, 140, 430, 240), (254, 243, 199), (217, 119, 6), "Semantic Classifier\n(Ollama AI)"),
        ((480, 140, 640, 240), (220, 252, 231), (22, 163, 74), "Safe Organizer\n(Collision Guard)"),
    ]

    for rect, fill_c, outline_c, text in box_coords:
        draw.rounded_rectangle(rect, radius=10, fill=fill_c, outline=outline_c, width=2)
        draw.text((rect[0] + 15, rect[1] + 35), text, fill=(30, 41, 59))

    # Connecting arrows
    draw.line([(220, 190), (270, 190)], fill=(148, 163, 184), width=3)
    draw.polygon([(265, 185), (270, 190), (265, 195)], fill=(148, 163, 184))

    draw.line([(430, 190), (480, 190)], fill=(148, 163, 184), width=3)
    draw.polygon([(475, 185), (480, 190), (475, 195)], fill=(148, 163, 184))

    # Footer note
    draw.text((60, 370), "LazySort Demo Diagram - Generated Visual Asset for Testing", fill=(100, 116, 139))

    image.save(str(filepath), format="PNG")


def generate_sample_dataset(target_dir: Path) -> List[Path]:
    """
    Populates target_dir with representative sample files across all categories.
    Returns the list of generated file paths.
    """
    dest = Path(target_dir)
    dest.mkdir(parents=True, exist_ok=True)
    created_files: List[Path] = []

    # 1. Finance - Invoice (PDF)
    f1 = dest / "Quarterly_Tax_Invoice_2026.pdf"
    create_sample_pdf(
        f1,
        "TAX INVOICE - CLOUD SERVICES INC.",
        [
            "Invoice Number: INV-2026-9041 | Date: October 15, 2026",
            "Billed To: Antigravity Labs Ltd. | VAT ID: US-994821049",
            "Itemized Services: High-Performance GPU Cloud Clusters (Compute hours: 140 hrs).",
            "Subtotal: $1,400.00 | Federal Tax (10%): $140.00",
            "TOTAL AMOUNT DUE: $1,540.00 USD. Payment terms: Net 30 via wire transfer.",
            "Please remit payment to Silicon Valley Banking Group, Routing #021000089."
        ]
    )
    created_files.append(f1)

    # 2. Work - Resume / CV (DOCX)
    f2 = dest / "Senior_Software_Engineer_Resume.docx"
    create_sample_docx(
        f2,
        "Alex Rivera - Senior Software Engineer",
        [
            ("Professional Summary", "Experienced backend engineer with 8+ years building distributed cloud applications, Python microservices, and AI integrations."),
            ("Work Experience", "Staff Engineer at ByteStream (2022-Present): Led the migration of data ingestion pipelines handling 5TB daily. Managed a team of 6 engineers."),
            ("Technical Skills", "Languages: Python, Go, TypeScript, SQL. Infrastructure: Docker, Kubernetes, AWS, Ollama, Redis, PostgreSQL."),
            ("Education", "Bachelor of Science in Computer Science, University of California, Berkeley.")
        ]
    )
    created_files.append(f2)

    # 3. Education - Lecture Notes (TXT)
    f3 = dest / "Algorithms_Lecture_04_Notes.txt"
    f3_text = (
        "CS 201: Data Structures & Algorithms - Lecture 04\n"
        "Topic: Divide-and-Conquer Sorting Algorithms\n"
        "Professor: Dr. Henderson | Midterm Date: November 12\n\n"
        "Key Concepts Covered:\n"
        "1. MergeSort: Recursively divides array into two halves, sorts each half, and merges them.\n"
        "   - Time Complexity: O(n log n) in best, average, and worst cases.\n"
        "   - Space Complexity: O(n) auxiliary space required for merging.\n\n"
        "2. QuickSort: Chooses a pivot element and partitions elements smaller/larger than pivot.\n"
        "   - Average Time Complexity: O(n log n), Worst Case: O(n^2) when poorly chosen pivot.\n\n"
        "Homework Assignment: Implement 3-way partitioning for duplicate elements by next Tuesday."
    )
    with open(f3, "w", encoding="utf-8") as f:
        f.write(f3_text)
    created_files.append(f3)

    # 4. Documents - Lease Agreement (PDF)
    f4 = dest / "Commercial_Office_Lease_Agreement.pdf"
    create_sample_pdf(
        f4,
        "STANDARD COMMERCIAL LEASE AGREEMENT",
        [
            "This Commercial Lease Agreement (the 'Agreement') is entered into as of October 1, 2026.",
            "BETWEEN: Metro Property Holdings LLC (the 'Landlord'), and Frontier AI Solutions (the 'Tenant').",
            "Premises: Suite 402, 550 Technology Boulevard, Austin, TX 78701.",
            "Term: The initial lease term shall commence on November 1, 2026, for a duration of 36 months.",
            "Security Deposit: Tenant agrees to deposit $8,500.00 prior to occupancy as security collateral.",
            "Governing Law: This agreement shall be governed by the laws of the State of Texas.",
            "IN WITNESS WHEREOF, the parties hereto have executed this Agreement on the day first written above."
        ]
    )
    created_files.append(f4)

    # 5. Images - Architecture Diagram (PNG)
    f5 = dest / "System_Architecture_Diagram.png"
    create_sample_png(f5, "LazySort AI Architecture & Data Flow")
    created_files.append(f5)

    # 6. Finance - Retail Coffee Receipt (TXT)
    f6 = dest / "Coffee_Shop_Receipt.txt"
    f6_text = (
        "THE ARTISAN ROASTER\n"
        "742 Evergreen Terrace\n"
        "Receipt #: 40992\n"
        "Date: 2026-10-22  09:14 AM\n"
        "-------------------------------\n"
        "1x Oat Milk Flat White     $5.50\n"
        "1x Almond Croissant        $4.75\n"
        "1x Avocado Toast           $9.00\n"
        "-------------------------------\n"
        "Subtotal:                 $19.25\n"
        "Sales Tax (8.25%):         $1.59\n"
        "Tip (20%):                 $3.85\n"
        "TOTAL PAID:               $24.69\n"
        "Payment Method: VISA ending in 8841\n"
        "Thank you for supporting local coffee!"
    )
    with open(f6, "w", encoding="utf-8") as f:
        f.write(f6_text)
    created_files.append(f6)

    # 7. Education - Course Syllabus (DOCX)
    f7 = dest / "Machine_Learning_Course_Syllabus.docx"
    create_sample_docx(
        f7,
        "CS 480: Introduction to Machine Learning - Fall Syllabus",
        [
            ("Course Overview", "An introduction to foundational supervised and unsupervised learning algorithms, gradient descent, neural networks, and model evaluation."),
            ("Grading Policy", "Homework assignments: 30%, Midterm Exam: 25%, Final Project: 35%, Participation: 10%."),
            ("Weekly Schedule", "Week 1: Linear Regression. Week 2: Logistic Regression & Softmax. Week 3: SVMs & Kernels. Week 4: Neural Networks. Week 5: Transformer Architectures."),
            ("Academic Integrity", "Students are expected to adhere to university standards. All submitted code must be original.")
        ]
    )
    created_files.append(f7)

    # 8. Others - App Config (JSON)
    f8 = dest / "application_runtime_config.json"
    f8_text = (
        "{\n"
        '  "app_name": "TelemetryDaemon",\n'
        '  "version": "2.4.1",\n'
        '  "port": 8080,\n'
        '  "metrics_enabled": true,\n'
        '  "retry_attempts": 5,\n'
        '  "timeout_seconds": 30,\n'
        '  "environment": "production"\n'
        "}\n"
    )
    with open(f8, "w", encoding="utf-8") as f:
        f.write(f8_text)
    created_files.append(f8)

    return created_files
